"""
Convert docs/corp-inventory/source/nexus-harbor.xlsx into per-category CSV files.
Run from repo root: python scripts/convert_inventory.py
Requires: openpyxl  (pip install openpyxl  or  uv run --with openpyxl python scripts/convert_inventory.py)
"""
import csv
import os
import sys

try:
    import openpyxl
except ImportError:
    print("openpyxl not found. Run: pip install openpyxl")
    sys.exit(1)

XLSX_PATH = "docs/corp-inventory/source/nexus-harbor.xlsx"
OUT_DIR = "docs/corp-inventory"

NPM_REPOS = {"npmjs", "npm-mts-design", "gitlab-npm"}
ALPINE_REPOS = {"alpine-proxy"}


def main():
    if not os.path.exists(XLSX_PATH):
        print(f"Not found: {XLSX_PATH}")
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)

    # --- Harbor ---
    ws = wb["Harbor"]
    harbor_rows = []
    for row in ws.iter_rows(values_only=True):
        project, repo = row[0], row[1]
        if project == "Проект" or not project or not repo:
            continue
        harbor_rows.append({"project": project, "repository": repo, "pull_path": f"harbor.mgts.ru/{repo}"})

    _write_csv(f"{OUT_DIR}/harbor-images.csv", ["project", "repository", "pull_path"], harbor_rows)
    print(f"harbor-images.csv: {len(harbor_rows)} images")

    # --- Nexus ---
    ws = wb["Nexus"]
    npm_rows, alpine_rows, pypi_rows, other_rows = [], [], [], []

    for row in ws.iter_rows(values_only=True):
        repo, url, val = row[0], row[1], row[2]
        if not repo or repo == "Repo" or not val:
            continue
        if repo in NPM_REPOS:
            npm_rows.append({"repo": repo, "url": url or "", "name": val})
        elif repo in ALPINE_REPOS:
            alpine_rows.append({"repo": repo, "version": val})
        elif "pypi" in repo.lower() or "python-proxy" in repo.lower():
            pypi_rows.append({"repo": repo, "url": url or "", "name": val})
        else:
            other_rows.append({"repo": repo, "url": url or "", "name": val})

    _write_csv(f"{OUT_DIR}/nexus-npm.csv", ["repo", "url", "name"], npm_rows)
    print(f"nexus-npm.csv: {len(npm_rows)} entries")

    _write_csv(f"{OUT_DIR}/nexus-alpine.csv", ["repo", "version"], alpine_rows)
    print(f"nexus-alpine.csv: {len(alpine_rows)} entries")

    if pypi_rows:
        _write_csv(f"{OUT_DIR}/nexus-pypi.csv", ["repo", "url", "name"], pypi_rows)
        print(f"nexus-pypi.csv: {len(pypi_rows)} entries")
    else:
        print("nexus-pypi.csv: no PyPI repo found (Python packages install from pypi.org via corp proxy)")

    _write_csv(f"{OUT_DIR}/nexus-other.csv", ["repo", "url", "name"], other_rows)
    print(f"nexus-other.csv: {len(other_rows)} entries")

    print("\nDone. Commit updated CSVs.")


def _write_csv(path, fieldnames, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)


if __name__ == "__main__":
    main()
